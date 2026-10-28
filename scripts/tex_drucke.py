#!/usr/bin/env python3
"""Liest Angaben zu den Drucken aus dem Druckeverzeichnis (texDateien/) aus.

Aufruf:  python scripts/tex_drucke.py

Je Druck werden Beschreibung, Titel, Stimmbuecher und die Verweise auf andere
Verzeichnisse (eine Spalte je Verzeichnis) ins Blatt Drucke von
raw/Drucke.xlsx geschrieben. Aus der RISM-Nummer des Verweises wird ueber die
Suche von RISM Online der Link zum Druck ermittelt. Nur leere Zellen werden
gefuellt, damit Korrekturen von Hand erhalten bleiben; Abweichungen werden nur
gemeldet. Siehe docs/entscheidungen.md, Abschnitt 26.
"""

from __future__ import annotations

import json
import re
import time
import unicodedata
import urllib.parse
import urllib.request
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
TEX_DIR = ROOT / "texDateien"
XLSX_PATH = ROOT / "raw" / "Drucke.xlsx"
SHEET = "Drucke"
# Die TeX-Dateien stammen aus DOS und sind in Codepage 437 kodiert.
ENCODING = "cp437"

# Verzeichnisse, auf die das Druckeverzeichnis verweist, in der Reihenfolge der
# Spalten. Verweise auf andere Verzeichnisse und Bemerkungen landen in
# "Verweise sonstige".
CATALOGUES = [
    "RISM", "Boetticher", "GA²", "GA", "EitnerQL", "EitnerV", "EitnerSW",
    "Delmotte/Dehn", "Huys", "HuysS", "Lesure/ThibaultRB", "Lesure/Thibault",
    "Lesure/ThibaultCh", "Lesure/ThibaultSuppl.", "N.Vogel", "Vogel Lasso",
    "Vogel-Einstein", "Göhler", "NR", "Fétis", "Mitjana/Davidsson",
    "Schmieder", "Lasso Motets", "Vanhulst", "Brown", "Meissner", "Cusick",
    "Forney", "Weaver",
]
# Schreibvarianten der Vorlage
CATALOGUE_ALIASES = {"Fètis": "Fétis", "Eitner QL": "EitnerQL"}
OTHER_REFERENCES = "sonstige"
REFERENCE_PREFIX = "Verweis "
RISM_LINK_COLUMN = "RISM"

COLUMNS = (
    ["Beschreibung", "Titel", "Stimmbücher"]
    + [REFERENCE_PREFIX + c for c in CATALOGUES + [OTHER_REFERENCES]]
    + [RISM_LINK_COLUMN]
)

# RISM-Nummern im Verweis: A/I als "[L 924" (Nachtrag "[LL 971 a"), B/I als
# hochgestellte Zaehlung hinter dem Jahr, z.B. "1565⁸"
RE_RISM_AI = re.compile(r"\[(LL?)\s*(\d+)(?:\s*([a-z])\b)?")
RE_RISM_BI = re.compile(r"\b(\d{4})([⁰¹²³⁴⁵⁶⁷⁸⁹]+)")
# Nur Drucke: Handschriften nennen oft die RISM-Nummer ihrer Druckvorlage
RISM_SEARCH = "https://rism.online/search?mode=sources&rows=20&fq=source-type:printed&q="

# Kopfzeile eines Drucks, z.B. "\ua 1581-5\ue", auch "ca. 1558-3" und in einem
# Fall in Spitzklammern
RE_HEADING = re.compile(r"\\ua\s*(.*?)\\ue", re.S)
RE_SIGLUM = re.compile(r"^(?:ca\. )?(?:\$\\langle\$)?(\d{4}-\d+(?:/[IVX]+)?)(?:\$\\rangle\$)?$")
# Nach Titel und Stimmbuechern beginnt die Beschreibung des Exemplars. Eine
# Tabelle (\settabs) kann auch im Titelblatt stehen und beendet ihn nicht.
RE_HEAD_END = re.compile(r"\\fg\s|Inhalt und|Inhalt entspricht")
RE_PARTBOOKS = re.compile(r"^(?:\\tg\s*)?Stimmb(?:ücher|uch)\s*:(.*)", re.S)
# Satzbefehle ohne Inhalt, z.B. \hyphenpenalty=50 oder \def\kopfzeilea{...}
RE_LAYOUT = re.compile(r"\\[a-zA-Z]+\s*=\s*-?[\d.]+\s*(?:pt|cm)?|\\def\\\w+\{[^{}]*\}")

SUPERSCRIPT = str.maketrans("0123456789o", "⁰¹²³⁴⁵⁶⁷⁸⁹°")
FROM_SUPERSCRIPT = str.maketrans("⁰¹²³⁴⁵⁶⁷⁸⁹", "0123456789")

# Ersetzungen vor dem Entfernen der uebrigen Befehle. Reihenfolge zaehlt.
REPLACEMENTS = [
    # Tabelle im Titelblatt: Musterzeile weg, Spalten und Zeilen als Leerraum
    (r"\\settabs.*?\\cr", " "),
    (r"\\\+|\\cr\b|(?<!\\)&", " "),
    (r"\$\\vert\$", " | "),
    (r"\$\\diamondsuit\$", " ◊ "),
    (r"\$\\langle\$", "⟨"),
    (r"\$\\rangle\$", "⟩"),
    (r"\$\\dagger\$", "†"),
    # Nasalstrich, z.B. Qv$\overline{\rm a}$do
    (r"\$\\overline\{\\rm (\w)\}\$", "\\1\u0304"),
    # uebergeschriebenes e, z.B. glu\llap{$^e$}ck
    (r"\\llap\{\$\^e\$\}", "\u0364"),
    (r"(?:\\hskip[\d.]+pt)?\\llap\{\$\^-\$\}", "\u0304"),
    # Akzente, z.B. BARR\`E, ru\"e, vtrius\'q;
    (r"\\`(\w)", "\\1\u0300"),
    (r"\\'(\w)", "\\1\u0301"),
    (r'\\"(\w)', "\\1\u0308"),
    # Umlaut als "u, in Fraktur als *u (dort uebergeschriebenes e)
    (r'"([aouAOU])', lambda m: {"a": "ä", "o": "ö", "u": "ü", "A": "Ä", "O": "Ö", "U": "Ü"}[m.group(1)]),
    (r"\*([aouAOU])", "\\1\u0364"),
    (r"\\oe\b\s?", "œ"),
    (r"\\OE\b\s?", "Œ"),
    (r"\\ae\b\s?", "æ"),
    (r"\\AE\b\s?", "Æ"),
    (r"\\ss\b\s?", "ß"),
    (r"\\etc\b", "etc."),
    (r"\\&", "&"),
    (r"\\-", ""),
    (r",,", "„"),
    (r"``", "“"),
    (r"~", " "),
    # hochgestellt, z.B. GA\ho 2 \ II oder 1565\ho 8 \
    (r"\\ho\s*(\w+)\s*\\?", lambda m: m.group(1).translate(SUPERSCRIPT) + " "),
    (r"\\[a-zA-Z]+\s?", ""),
    (r"\\\s", " "),
    (r"[{}$]", ""),
]


def clean_tex(text):
    for pattern, replacement in REPLACEMENTS:
        text = re.sub(pattern, replacement, text)
    return unicodedata.normalize("NFC", re.sub(r"\s+", " ", text).strip())


def split_blocks(text):
    """Zerlegt eine TeX-Datei in Abschnitte, je einer pro Druck."""
    starts = []
    for match in RE_HEADING.finditer(text):
        siglum = RE_SIGLUM.match(match.group(1).strip())
        if siglum:
            starts.append((match.start(), siglum.group(1)))
    for (start, siglum), following in zip(starts, starts[1:] + [(len(text), None)]):
        yield siglum, text[start:following[0]]


def catalogue_of(reference):
    """Liefert (Verzeichnis, Angabe) eines Verweises oder (None, Verweis)."""
    for name in sorted(list(CATALOGUES) + list(CATALOGUE_ALIASES), key=len, reverse=True):
        rest = reference[len(name):]
        if reference.startswith(name) and not rest[:1].isalpha():
            return CATALOGUE_ALIASES.get(name, name), rest.strip()
    return None, reference


def split_references(references):
    """Verteilt die Verweise auf eine Spalte je Verzeichnis.

    Ein Verweis kann mehrere Verzeichnisse nennen, getrennt durch
    Semikolon ("EitnerV 1555; EitnerQL 1555"); getrennt wird nur, wenn der
    Teil nach dem Semikolon mit einem bekannten Verzeichnis beginnt.
    """
    columns = {}
    for reference in references:
        parts = [reference]
        while True:
            head, sep, tail = parts[-1].partition("; ")
            if not sep or catalogue_of(tail)[0] is None:
                break
            parts[-1:] = [head, tail]
        for part in parts:
            name, value = catalogue_of(part)
            column = REFERENCE_PREFIX + (name if name and value else OTHER_REFERENCES)
            if not (name and value):
                value = part
            columns[column] = "; ".join(filter(None, [columns.get(column), value]))
    return columns


def rism_numbers(reference):
    """Liefert die RISM-Nummern eines Verweises, A/I vor B/I."""
    numbers = ["A/I {0} {1}{2}".format(*m) for m in RE_RISM_AI.findall(reference)]
    for year, count in RE_RISM_BI.findall(reference):
        numbers.append("B/I {0}/{1}".format(year, count.translate(FROM_SUPERSCRIPT)))
    return numbers


def search_rism(number):
    """Liefert die Treffer der Phrasensuche, bei Netzfehlern bis zu drei Versuche."""
    request = urllib.request.Request(
        RISM_SEARCH + urllib.parse.quote('"{0}"'.format(number)),
        headers={"Accept": "application/ld+json"},
    )
    for attempt in range(3):
        try:
            with urllib.request.urlopen(request, timeout=60) as response:
                return json.load(response).get("items", [])
        except OSError:
            time.sleep(5 * (attempt + 1))
    raise OSError("RISM Online nicht erreichbar: " + number)


def find_rism_link(numbers):
    """Sucht die Nummern bei RISM Online, liefert den Link bei genau einem Treffer."""
    for number in numbers:
        items = search_rism(number)
        time.sleep(0.5)
        if len(items) == 1:
            return items[0]["id"]
    return None


def parse_block(block):
    """Liefert Beschreibung, Titel, Stimmbuecher und Verweise eines Drucks."""
    head = RE_LAYOUT.sub("", RE_HEAD_END.split(block, maxsplit=1)[0])
    # Erste Ueberschrift ist die Sigle, die weiteren die Beschreibung
    headings = [m.group(1) for m in RE_HEADING.finditer(head)][1:]
    record = {
        "Beschreibung": clean_tex(" ".join(headings)) or None,
        "Verweise": None,
        "Titel": None,
        "Stimmbücher": None,
    }
    for segment in re.split(r"\\medskip|\\bigskip|\\smallskip", RE_HEADING.sub("", head)):
        segment = segment.strip()
        if not segment or re.match(r"(?:\\tg\s*)?LV ", segment):
            continue
        partbooks = RE_PARTBOOKS.match(segment)
        if partbooks:
            record["Stimmbücher"] = clean_tex(partbooks.group(1))
        elif segment == "Chorbuch":
            record["Stimmbücher"] = segment
        elif record["Stimmbücher"] is not None:
            # danach folgt nur noch das Inhaltsverzeichnis
            break
        elif record["Verweise"] is None and (
            "diamondsuit" in segment or segment.startswith("RISM")
            or (record["Titel"] is None and "$\\vert$" not in segment)
        ):
            references = [clean_tex(r).strip("; ") for r in segment.split("$\\diamondsuit$")]
            record["Verweise"] = [r for r in references if r]
        elif record["Titel"] is None:
            record["Titel"] = clean_tex(segment)
    record.update(split_references(record.pop("Verweise") or []))
    return record


def read_tex():
    """Liest alle TeX-Dateien in der Reihenfolge LVtex1 bis LVTEX9.

    Steht ein Druck in zwei Dateien (1578-1 bis 1578-9 in LVtex4 und
    LVTEX5), gilt der Abschnitt aus der spaeteren Datei.
    """
    prints = {}
    duplicates = []
    for path in sorted(TEX_DIR.iterdir(), key=lambda p: p.name.lower()):
        text = path.read_bytes().decode(ENCODING)
        for siglum, block in split_blocks(text):
            if siglum in prints:
                duplicates.append(siglum)
            prints[siglum] = parse_block(block)
    return prints, duplicates


def main():
    prints, duplicates = read_tex()
    print("{0} Drucke im Druckeverzeichnis".format(len(prints)))

    workbook = openpyxl.load_workbook(XLSX_PATH)
    sheet = workbook[SHEET]
    header = [cell.value for cell in sheet[1]]
    for name in COLUMNS:
        if name not in header:
            header.append(name)
            sheet.cell(row=1, column=len(header), value=name)
    column_of = {name: header.index(name) + 1 for name in COLUMNS}
    siglum_column = header.index("Sigle") + 1

    filled = 0
    conflicts = []
    seen = set()
    missing_in_tex = []
    from_whole_print = []
    rism_cache = {}
    rism_not_found = []
    rism_failed = []
    for row in range(2, sheet.max_row + 1):
        siglum = sheet.cell(row=row, column=siglum_column).value
        if not siglum:
            continue
        # Teile eines mehrbaendigen Drucks (1575-3/I) beschreibt das
        # Verzeichnis nur gemeinsam unter der Grundsigle
        source = siglum if siglum in prints else siglum.split("/")[0]
        record = prints.get(source)
        if record is None:
            missing_in_tex.append(siglum)
            continue
        if source != siglum:
            from_whole_print.append(siglum)
        seen.add(source)
        record = dict(record)
        # RISM Online nur abfragen, solange die Zelle leer ist
        numbers = rism_numbers(record.get(REFERENCE_PREFIX + "RISM") or "")
        if numbers and sheet.cell(row=row, column=column_of[RISM_LINK_COLUMN]).value in (None, ""):
            key = tuple(numbers)
            if key not in rism_cache:
                try:
                    rism_cache[key] = find_rism_link(numbers)
                except OSError:
                    # Zelle bleibt leer und wird beim naechsten Lauf erneut abgefragt
                    rism_failed.append(siglum)
                    rism_cache[key] = None
            record[RISM_LINK_COLUMN] = rism_cache[key]
            if rism_cache[key] is None and siglum not in rism_failed:
                rism_not_found.append("{0} ({1})".format(siglum, ", ".join(numbers)))
        for name, value in record.items():
            cell = sheet.cell(row=row, column=column_of[name])
            if value is None:
                continue
            if cell.value in (None, ""):
                cell.value = value
                filled += 1
            elif cell.value != value:
                conflicts.append((siglum, name))
    workbook.save(XLSX_PATH)

    missing_fields = {
        name: sorted(s for s in seen if prints[s].get(name) is None)
        for name in ["Beschreibung", "Titel", "Stimmbücher", REFERENCE_PREFIX + "RISM"]
    }
    print("{0} Zellen gefuellt, {1} abweichende Zellen nicht ueberschrieben".format(
        filled, len(conflicts)))
    for siglum, name in conflicts:
        print("  abweichend: {0} / {1}".format(siglum, name))
    print("Angaben des Gesamtdrucks uebernommen: {0}".format(
        ", ".join(from_whole_print) or "-"))
    print("Nur in Drucke.xlsx: {0}".format(", ".join(missing_in_tex) or "-"))
    print("Nur im Druckeverzeichnis: {0}".format(
        ", ".join(sorted(set(prints) - seen)) or "-"))
    print("Doppelt im Druckeverzeichnis (spaeterer Abschnitt gilt): {0}".format(
        ", ".join(duplicates) or "-"))
    for name, sigla in missing_fields.items():
        print("Ohne {0}: {1} ({2})".format(name, len(sigla), ", ".join(sigla[:15])))
    print("RISM-Nummer ohne eindeutigen Treffer bei RISM Online: {0}".format(len(rism_not_found)))
    for item in rism_not_found:
        print("  " + item)
    print("RISM Online nicht erreicht, bitte erneut ausfuehren: {0}".format(
        ", ".join(rism_failed) or "-"))


if __name__ == "__main__":
    main()
